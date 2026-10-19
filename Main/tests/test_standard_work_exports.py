"""Regression checks for the imported-template SWCS and annotated SWS."""
import re
import copy
import unittest
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from standard_work_exports import planned_steps, swip_total, export_sws_pptx, export_swcs_xlsx

NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def xml_cell(root, coordinate):
    cell=next((x for x in root.findall('.//m:sheetData/m:row/m:c',NS) if x.get('r')==coordinate),None)
    if cell is None:return None
    if cell.get('t')=='inlineStr':
        return ''.join(x.text or '' for x in cell.findall('.//m:t',NS))
    v=cell.find('m:v',NS)
    return None if v is None else v.text


class StandardWorkExportsTest(unittest.TestCase):
    @staticmethod
    def sample():
        return {
            'takt_time':12.0,'layout_scale_px_per_m':50,'walking_speed_m_per_s':1.2,
            'line_name':'Example line',
            'stations':{
                '1':{'name':'Inspect','x':100,'y':100,'w':150,'h':96,'has_part':True,
                     'processes':[{'id':'p1','name':'Manual inspection','manual_time':2.,'auto_time':1.}]},
                '2':{'name':'Polish','x':350,'y':100,'w':150,'h':96,'has_part':False,
                     'processes':[{'id':'p2','name':'Polish surface','manual_time':3.,'auto_time':0.}]},
            },
            'boxes':{
                'in1':{'type':'in','name':'Raw','x':5,'y':100,'w':50,'h':55,'start_count':7},
                'out1':{'type':'out','name':'Finished','x':580,'y':100,'w':50,'h':55,'capacity':20},
                'swip1':{'type':'swip','name':'Buffer','x':300,'y':300,'w':50,'h':55,'has_part':True},
                'dir':{'type':'direction','name':'Arrow','x':0,'y':200,'w':600,'h':30},
            },
            'operators':{
                'Op2':{'name':'Op2','x':400,'y':250,'color':[30,105,180],
                       'sequence':[{'action':'work','loc':'2','process':'p2'}]},
                'Op1':{'name':'Op1','x':175,'y':74,'color':[210,60,50],
                       'attached_to':{'station_id':'1','side':'top'},
                       'sequence':[{'action':'work','loc':'1','process':'p1','standing_side':'top'},
                                   {'action':'pick','loc':'1'},{'action':'drop','loc':'2'}]},
            },
        }

    def test_time_values_and_category_counts(self):
        scenario=self.sample()
        self.assertEqual(swip_total(scenario),9)
        rows=planned_steps(scenario,'Op1')
        self.assertEqual(rows[0]['manual'],2)
        self.assertEqual(rows[0]['auto'],1)
        self.assertEqual(rows[1]['action'],'pick')
        self.assertGreater(rows[2]['walk'],0)
        self.assertAlmostEqual(rows[2]['walk'],rows[2]['distance_m']/1.2/60,places=6)

    def test_sws_has_part_icons_dotted_boxes_arrows_and_circular_numbers(self):
        p=Presentation(BytesIO(export_sws_pptx(self.sample()).getvalue()))
        self.assertEqual(len(p.slides),3)
        for slide in p.slides:
            pictures=[sh for sh in slide.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE]
            self.assertGreaterEqual(len(pictures),4)  # arrow + starting Part icons; no injected legend icon
            text=' '.join(sh.text for sh in slide.shapes if sh.has_text_frame)
            self.assertIn('SWIP',text)
            self.assertNotIn('SWIP: Buffer',text)
            self.assertNotIn('Part icon = starting SWIP',text)  # supplied legend left untouched
            for label in ('IN: Raw','OUT: Finished','SWIP'):
                matches=[sh for sh in slide.shapes if sh.has_text_frame and label in sh.text]
                self.assertTrue(matches,label)
                self.assertTrue(any(sh.line.dash_style==MSO_LINE_DASH_STYLE.ROUND_DOT for sh in matches),label)
        op_slide=p.slides[1]
        numbered=[sh for sh in op_slide.shapes if sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE and
                  sh.auto_shape_type.name=='OVAL' and sh.has_text_frame and sh.text.strip().isdigit()]
        numbers={sh.text.strip() for sh in numbered}
        self.assertTrue({'1','2','3'}.issubset(numbers))
        self.assertTrue(any(sh.shape_type==MSO_SHAPE_TYPE.FREEFORM and
                            sh.fill.type is not None for sh in op_slide.shapes))

    def test_sws_template_frame_is_unchanged_and_canvas_is_separate(self):
        from copy import deepcopy
        from pathlib import Path
        base=Presentation(str(Path(__file__).resolve().parents[1] / 'resources' / 'SWS_Frame.pptx'))
        output=Presentation(BytesIO(export_sws_pptx(self.sample()).getvalue()))
        original=[sh._element.xml for sh in base.slides[0].shapes]
        self.assertEqual(len(output.slides),3)
        for slide in output.slides:
            self.assertEqual([sh._element.xml for sh in list(slide.shapes)[:len(original)]],original)
            # Generated additions must remain inside the blank drawing region;
            # labels, operator movement lines and part icons must not cross
            # the user-editable header, legend or footer regions.
            for sh in list(slide.shapes)[len(original):]:
                self.assertGreaterEqual(sh.top / 914400, 1.15, sh.name)
                self.assertLessEqual((sh.top+sh.height) / 914400, 6.17, sh.name)
            # The frame's blank user-editable data fields remain blank.
            tables=[sh.table for sh in slide.shapes if sh.has_table]
            self.assertEqual(tables[0].cell(1,3).text,'')
            self.assertEqual(tables[1].cell(0,1).text,'')

    def test_sws_preserves_direction_transform(self):
        scenario=self.sample()
        scenario['boxes']['dir'].update({'rotation':90,'flip_horizontal':True,'flip_vertical':True})
        p=Presentation(BytesIO(export_sws_pptx(scenario).getvalue()))
        for slide in p.slides:
            pictures=[x for x in slide.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE]
            arrows=[x for x in pictures if abs(x.rotation-90)<0.001]
            self.assertEqual(len(arrows),1)
            self.assertEqual(arrows[0]._element.spPr.xfrm.get('flipH'),'1')
            self.assertEqual(arrows[0]._element.spPr.xfrm.get('flipV'),'1')

    def test_swcs_actual_template_preserves_drawings_and_blue_fields(self):
        buf=export_swcs_xlsx(self.sample())
        with ZipFile(buf) as z:
            self.assertEqual(len([n for n in z.namelist() if re.fullmatch(r'xl/drawings/drawing\d+\.xml',n)]),2)
            sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(xml_cell(sheet,'CN3'),'OP1')
            self.assertEqual(xml_cell(sheet,'D12'),'Work: Manual inspection — S1 — Inspect')
            self.assertEqual(xml_cell(sheet,'G12'),'2 min')
            self.assertEqual(xml_cell(sheet,'H12'),'1 min')
            self.assertNotIn('Seal removal',z.read('xl/worksheets/sheet1.xml').decode())
            # Editable factory drawing objects have not been discarded.
            root=ET.fromstring(z.read('xl/drawings/drawing1.xml'))
            self.assertGreater(len(root),5)
            self.assertIn('drawing1.xml',z.read('xl/worksheets/_rels/sheet1.xml.rels').decode())

    def test_swcs_continuation_and_many_operators(self):
        scenario=self.sample()
        for i in range(3,9): scenario['operators'][f'Op{i}']=copy.deepcopy(scenario['operators']['Op1'])
        scenario['operators']['Op1']['sequence']*=6
        buf=export_swcs_xlsx(scenario)
        with ZipFile(buf) as z:
            book=ET.fromstring(z.read('xl/workbook.xml'))
            names=[x.get('name') for x in book.findall('.//m:sheet',NS)]
            self.assertEqual(len(names),9)
            self.assertIn('Op1 (2)',names)
            self.assertIn('Op8',names)
            cont=ET.fromstring(z.read('xl/worksheets/sheet2.xml'))
            self.assertEqual(xml_cell(cont,'C12'),'11')

if __name__=='__main__':unittest.main()
