"""Read the reference-style PQPR workbook into an editable project.

No formulas are evaluated. P/X applicability and process headers are read;
manual/automatic timings are left unknown because the reference has none.
"""
import colorsys
from io import BytesIO
from pathlib import Path
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from scenario_schema import ScenarioValidationError
from pqpr_schema import validate_pqpr


def clean(value):
    return ' '.join(str(value).split()) if value is not None else ''


def header_layout(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50)):
        values = {re.sub(r'[^a-z]', '', clean(cell.value).lower()): cell.column for cell in row if cell.value is not None}
        frame_col = values.get('frame') or values.get('framesize')
        end_meta = values.get('totalpieces')
        if frame_col and end_meta:
            columns = [col for col in range(end_meta + 1, sheet.max_column + 1) if clean(sheet.cell(row[0].row, col).value)]
            if columns:
                return row[0].row, frame_col, columns, values
    return None


def workbook_from_bytes(raw):
    try:
        return load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise ScenarioValidationError(['Cannot read this Excel file. Choose an .xlsx PQPR workbook.']) from exc


def worksheets(raw):
    workbook = workbook_from_bytes(raw)
    try:
        result = [{'name': sheet.title, 'process_count': len(header_layout(sheet)[2])}
                  for sheet in workbook if header_layout(sheet)]
        if not result:
            raise ScenarioValidationError(['No PQPR worksheet found. Expected Frame/Framesize and Total Pieces headers followed by process columns.'])
        return result
    finally:
        workbook.close()


def color_category(cell, workbook):
    color = cell.fill.fgColor
    rgb = None
    if color.type == 'rgb':
        rgb = color.rgb[-6:]
    elif color.type == 'indexed' and color.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[color.indexed][-6:]
    elif color.type == 'theme' and workbook.loaded_theme:
        try:
            root = ET.fromstring(workbook.loaded_theme)
            scheme = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
            item = list(list(scheme)[color.theme])[0]
            rgb = item.attrib.get('lastClr') or item.attrib.get('val')
        except (ET.ParseError, IndexError, TypeError):
            pass
    if rgb and re.fullmatch(r'[0-9A-Fa-f]{6}', rgb):
        channels = [int(rgb[i:i+2], 16) / 255 for i in (0, 2, 4)]
        hue, saturation, _ = colorsys.rgb_to_hsv(*channels)
        hue *= 360
        if saturation > .2:
            if 65 <= hue < 175: return 'in_line'
            if 175 <= hue <= 280: return 'unique'
            if 10 <= hue < 50: return 'out_of_line'
    return None


def import_worksheet(raw, sheet_name, filename):
    workbook = workbook_from_bytes(raw)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ScenarioValidationError(['Choose an existing PQPR worksheet.'])
        sheet = workbook[sheet_name]
        layout = header_layout(sheet)
        if not layout:
            raise ScenarioValidationError(['This worksheet does not have the supported PQPR header layout.'])
        header_row, frame_col, columns, metadata = layout
        project = {'document_type': 'pqpr', 'schema_version': 1,
                   'project_name': Path(filename).stem + ' — ' + sheet_name,
                   'processes': [], 'framesizes': [],
                   'source': {'filename': filename, 'worksheet': sheet_name}}
        warnings = ['Manual and automatic times were not provided by this workbook layout. Enter them in the process list.']
        unknown_colors = []
        for index, col in enumerate(columns):
            cell = sheet.cell(header_row, col)
            category = color_category(cell, workbook)
            if category is None:
                category = 'out_of_line'
                unknown_colors.append(clean(cell.value))
            project['processes'].append({'id': f'p{index+1}', 'name': clean(cell.value),
                                        'category': category, 'manual_time': None, 'auto_time': None,
                                        'source_column': cell.column_letter})
        for row in range(header_row + 1, sheet.max_row + 1):
            name = clean(sheet.cell(row, frame_col).value)
            marks = [clean(sheet.cell(row, col).value).upper() for col in columns]
            if not name or not any(mark in ('P', 'X') for mark in marks):
                continue
            frame = {'id': f'f{len(project["framesizes"])+1}', 'name': name, 'stage': '', 'component': '', 'applicability': {}}
            for key in ('stage', 'component'):
                if metadata.get(key): frame[key] = clean(sheet.cell(row, metadata[key]).value)
            for process, mark in zip(project['processes'], marks):
                if mark in ('P', 'X'):
                    frame['applicability'][process['id']] = 'required' if mark == 'P' else 'optional'
            project['framesizes'].append(frame)
        if not project['framesizes']:
            raise ScenarioValidationError(['No framesize rows with P/X process marks were found.'])
        if unknown_colors:
            warnings.append('Unrecognized header colors were set to orange. Review: ' + ', '.join(unknown_colors) + '.')
        return validate_pqpr(project), warnings
    finally:
        workbook.close()
