"""Draft Standard Work Sheet (PowerPoint) and Combination Sheet (Excel) exports.

The F2 SWS and HA S4N SWCS provided with the project informed the format:
- a plant-layout sheet with operator work paths and summary fields;
- an operator-per-sheet combination chart for Manual, Automatic and Walking.
This exporter intentionally does not fabricate detailed work instructions, quality checks,
safety points, waiting time, or certification information absent from the scenario.
"""
from __future__ import annotations

from datetime import date
from copy import deepcopy
from io import BytesIO
import math
import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from walking_geometry import box_center, station_destination, walking_metrics

BLUE = '#193D65'
PALE = '#F2F6FB'
BORDER = '#75859A'
MANUAL = '#F2A065'
AUTO = '#73B6CA'
WALK = '#A697DD'
WAIT = '#C3C9D1'


def rgb(value):
    value = (value or '#000000').lstrip('#')
    return RGBColor.from_string(value[:6].upper())


def fill_rgb(value):
    return PatternFill('solid', fgColor=value.lstrip('#'))


def operator_ids(scenario):
    def sort_key(name):
        match = re.fullmatch(r'Op(\d+)', str(name), re.IGNORECASE)
        return (0, int(match.group(1)), '') if match else (1, 0, str(name).casefold())
    return sorted(scenario.get('operators', {}), key=sort_key)


def scenario_name(scenario):
    for field in ('line_name', 'scenario_name', 'name', 'title'):
        value = scenario.get(field)
        if isinstance(value, str) and value.strip():
            return value.strip()[:80]
    return 'Configured Production Line'


def swip_total(scenario):
    stations = sum(bool(s.get('has_part')) for s in scenario.get('stations', {}).values())
    swip_boxes = sum(bool(b.get('has_part')) for b in scenario.get('boxes', {}).values() if b.get('type') == 'swip')
    in_parts = sum(max(0, int(b.get('start_count') or 0)) for b in scenario.get('boxes', {}).values() if b.get('type') == 'in')
    return stations + swip_boxes + in_parts


def category_label(scenario, loc):
    if loc in scenario.get('stations', {}):
        return f"S{loc} — {scenario['stations'][loc].get('name', f'Station {loc}') }"
    box = scenario.get('boxes', {}).get(loc)
    if not box:
        return str(loc)
    return f"{box.get('name', loc)} ({box.get('type', '').upper()})"


def planned_steps(scenario, operator_id):
    """Return scenario-based first-pass times, not certified work instructions.

    Automatic dwell starts when manual work ends, and can overlap the operator's
    next step. Travel is straight-line at the configured physical scale/speed.
    """
    operator = scenario['operators'][operator_id]
    position = {'x': float(operator['x']), 'y': float(operator['y'])}
    home = dict(position)
    scale = scenario.get('layout_scale_px_per_m', 50.0)
    speed = scenario.get('walking_speed_m_per_s', 1.2)
    clock = 0.0
    rows = []
    stations = scenario['stations']
    boxes = scenario['boxes']
    for index, step in enumerate(operator.get('sequence', []), 1):
        loc = step.get('loc')
        if loc in stations:
            target = station_destination(stations, loc, step, position, operator)
        elif loc in boxes:
            target = box_center(boxes[loc])
        else:
            continue
        distance_m, walk_min = walking_metrics(position, target, scale, speed)
        action = step.get('action')
        destination_label = category_label(scenario, loc)
        manual = automatic = 0.0
        if action == 'work' and loc in stations:
            process = next((p for p in stations[loc].get('processes', []) if p.get('id') == step.get('process')), None)
            if process:
                manual = float(process.get('manual_time') or 0)
                automatic = float(process.get('auto_time') or 0)
                description = f"Work: {process.get('name', 'Process')} — {destination_label}"
            else:
                description = f"Work — {destination_label}"
        elif action == 'pick':
            description = f"Pick part from {destination_label}"
        elif action == 'drop':
            description = f"Drop part at {destination_label}"
        else:
            description = f"{action or 'Visit'} — {destination_label}"
        walk_start = clock
        manual_start = walk_start + walk_min
        auto_start = manual_start + manual
        rows.append({
            'index': index, 'action': action, 'description': description, 'loc': loc,
            'manual': manual, 'auto': automatic, 'walk': walk_min, 'distance_m': distance_m,
            'walk_start': walk_start, 'manual_start': manual_start,
            'auto_start': auto_start, 'auto_end': auto_start + automatic,
            'position': target,
        })
        clock = auto_start  # machine automatic time does not occupy the operator
        position = target
    if rows:
        distance_m, walk_min = walking_metrics(position, home, scale, speed)
        if walk_min > 1e-9:
            rows.append({'index': len(rows) + 1, 'action': 'return',
                         'description': 'Walk back to operator home position', 'loc': None,
                         'manual': 0.0, 'auto': 0.0, 'walk': walk_min, 'distance_m': distance_m,
                         'walk_start': clock, 'manual_start': clock + walk_min,
                         'auto_start': clock + walk_min, 'auto_end': clock + walk_min,
                         'position': home})
            clock += walk_min
    return rows


def sws_box(slide, x, y, w, h, text='', *, bg='#FFFFFF', color='#000000', border=BORDER, size=9, bold=False, dotted=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(bg)
    shape.line.color.rgb = rgb(border); shape.line.width = Pt(0.85 if dotted else 0.65)
    if dotted: shape.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.035);tf.margin_right = Inches(0.035)
    tf.margin_top = Inches(0.02);tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    run = para.add_run(); run.text = str(text)
    run.font.name = 'Arial'; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = rgb(color)
    return shape


def sws_text(slide, x, y, w, h, text, size=9, color='#26384C', bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.015)
    para = tf.paragraphs[0]; para.alignment = align
    run = para.add_run(); run.text = str(text)
    run.font.name = 'Arial'; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def route_color(operator):
    color = operator.get('color', [42, 111, 188])
    return '#%02X%02X%02X' % tuple(max(0, min(int(x), 255)) for x in color[:3])


def bezier_points(a, b, bend=0.12, segments=22):
    """Smooth quarter-bend between destinations, with a gentle offset."""
    dx, dy = b[0] - a[0], b[1] - a[1]
    dist = max(0.001, math.hypot(dx, dy))
    perp = (-dy / dist, dx / dist)
    bend = min(abs(bend), 0.28) * min(1.0, dist * 2.5)
    c1 = (a[0] + dx * 0.34 + perp[0] * bend, a[1] + dy * 0.34 + perp[1] * bend)
    c2 = (a[0] + dx * 0.67 + perp[0] * bend, a[1] + dy * 0.67 + perp[1] * bend)
    result = []
    for i in range(segments + 1):
        t = i / segments; q = 1 - t
        result.append((q**3*a[0] + 3*q*q*t*c1[0] + 3*q*t*t*c2[0] + t**3*b[0],
                       q**3*a[1] + 3*q*q*t*c1[1] + 3*q*t*t*c2[1] + t**3*b[1]))
    return result


def add_curved_path(slide, start, end, color, bend):
    """Editable curved route with a small filled arrow head at its destination."""
    if math.hypot(end[0] - start[0], end[1] - start[1]) < 0.03:
        return
    pts = bezier_points(start, end, bend, 18)
    builder = slide.shapes.build_freeform(int(Inches(pts[0][0])), int(Inches(pts[0][1])))
    builder.add_line_segments([(int(Inches(x)), int(Inches(y))) for x, y in pts[1:]], close=False)
    shape = builder.convert_to_shape()
    shape.fill.background(); shape.line.color.rgb = rgb(color); shape.line.width = Pt(2.15)
    # Triangle is an editable PowerPoint shape. Its tip points toward the
    # destination, not back toward the origin.
    tip = pts[-1]
    earlier = pts[-3]
    dx,dy = tip[0]-earlier[0],tip[1]-earlier[1]
    dist = math.hypot(dx,dy)
    if dist < .0001:
        return
    ux,uy=dx/dist,dy/dist
    tip=(tip[0]-.015*ux,tip[1]-.015*uy)
    base=(tip[0]-.13*ux,tip[1]-.13*uy)
    perp=(-uy*.065,ux*.065)
    vertices=[tip,(base[0]+perp[0],base[1]+perp[1]),(base[0]-perp[0],base[1]-perp[1])]
    head=slide.shapes.build_freeform(int(Inches(vertices[0][0])),int(Inches(vertices[0][1])))
    head.add_line_segments([(int(Inches(x)),int(Inches(y))) for x,y in vertices[1:]],close=True)
    head=head.convert_to_shape();head.fill.solid();head.fill.fore_color.rgb=rgb(color)
    head.line.fill.background()


def add_part_icon(slide, x, y, size=.18, count=None):
    """The same Part.png glyph used in the browser, on a contrasting disc."""
    bg=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    bg.fill.solid();bg.fill.fore_color.rgb=rgb('#1D3954');bg.line.fill.background()
    part=Path(__file__).resolve().parent/'static'/'images'/'Part.png'
    if part.is_file():
        slide.shapes.add_picture(str(part), Inches(x+.012), Inches(y+.012), Inches(size-.024), Inches(size-.024))
    if count is not None and count > 1:
        sws_text(slide,x+size*.65,y-.10,.43,.14,f'×{count}',size=7.4,color='#1D3954',bold=True)


def add_numbered_steps(slide, visits, *, reserved=()):
    """Position simple numbered stops, avoiding operator start badges and peers."""
    taken = [tuple(pos) for pos in reserved]
    for visit in sorted(visits, key=lambda item: (item['operator'], item['number'])):
        anchor = visit['point']
        hue = visit['hue']
        number = visit['number']
        d = .235
        best = None
        for ring in range(1, 12):
            radius = .24 + (ring - 1) * .10
            for angle_index in range(18):
                theta = (2 * math.pi * angle_index / 18) + (number % 5) * .18
                x = anchor[0] + radius * math.cos(theta) - d / 2
                y = anchor[1] + radius * math.sin(theta) - d / 2
                overflow = max(0, .25 - x) + max(0, x + d - 13.05) + max(0, 1.28 - y) + max(0, y + d - 6.04)
                collision = sum(max(0, d + .12 - math.hypot((x + d/2) - ox, (y + d/2) - oy))**2 for ox, oy in taken)
                score = overflow * 500 + collision * 800 + radius * .15
                if best is None or score < best[0]:
                    best = (score, x, y)
            if best and best[0] < .14:
                break
        _, x, y = best
        x = max(.25, min(13.05 - d, x))
        y = max(1.28, min(6.04 - d, y))
        taken.append((x + d/2, y + d/2))
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(anchor[0]), Inches(anchor[1]),
                                          Inches(x + d/2), Inches(y + d/2))
        line.line.color.rgb = rgb(hue)
        line.line.width = Pt(.6)
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        badge.fill.solid(); badge.fill.fore_color.rgb = rgb('#FFFFFF')
        badge.line.color.rgb = rgb(hue); badge.line.width = Pt(1.2)
        tf = badge.text_frame; tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
        run = para.add_run(); run.text = str(number)
        run.font.bold = True; run.font.size = Pt(7.7 if len(str(number)) < 3 else 6.6)
        run.font.color.rgb = rgb(hue)


def movement_visits(scenario, operator_id):
    """A *standard repeatable loop*, not a chronological event trace.

    - Each distinct assigned station/material location appears once.
    - Work/Pick/Drop at the same location do not create new route stops.
    - An initial pickup from IN/SWIP is shifted to the end of the loop, so the
      route starts at the first workstation and shows where the operator
      collects material *before the next cycle*.
    - Home/idle position is not a route stop. The slide closes back to the
      first station instead. This has NO effect on SimPy walking calculations.
    """
    operator = scenario.get('operators', {}).get(operator_id, {})
    stations = scenario.get('stations', {})
    boxes = scenario.get('boxes', {})
    unique = []
    seen = set()
    last_position = {'x': float(operator.get('x', 0)), 'y': float(operator.get('y', 0))}
    for step in operator.get('sequence', []):
        loc = str(step.get('loc', ''))
        if loc not in stations and loc not in boxes:
            continue
        if loc in seen:
            continue
        if loc in stations:
            target = station_destination(stations, loc, step, last_position, operator)
            kind = 'station'
        elif boxes[loc].get('type') in ('swip', 'in', 'out'):
            target = box_center(boxes[loc])
            kind = boxes[loc].get('type')
        else:
            continue
        seen.add(loc)
        unique.append({'loc': loc, 'kind': kind,
                       'point': (float(target['x']), float(target['y']))})
        last_position = target

    if any(item['kind'] == 'station' for item in unique):
        # Common cycle: SWIP -> first work station -> remaining stations.
        # Show instead: first station -> ... -> SWIP -> first station.
        first_station = next(i for i, item in enumerate(unique) if item['kind'] == 'station')
        if first_station and all(item['kind'] in ('swip', 'in') for item in unique[:first_station]):
            unique = unique[first_station:] + unique[:first_station]

    for index, item in enumerate(unique, 1):
        item['number'] = index
    return unique


def add_operator_start_icon(slide, operator_id, hue, point, offset_index=0):
    """Builder-style operator circle at the first assigned station."""
    d = .32
    # Keep multiple operators using the same station distinguishable.
    angle = math.radians(-50 + offset_index * 52)
    radius = .13 if offset_index else 0
    cx = max(.45, min(12.90, point[0] + math.cos(angle) * radius))
    cy = max(1.50, min(5.85, point[1] + math.sin(angle) * radius))
    if math.dist((cx, cy), point) > .025:
        tether = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(point[0]), Inches(point[1]),
                                            Inches(cx), Inches(cy))
        tether.line.color.rgb = rgb(hue); tether.line.width = Pt(.75)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                   Inches(cx - d/2), Inches(cy - d/2), Inches(d), Inches(d))
    badge.fill.solid(); badge.fill.fore_color.rgb = rgb(hue)
    badge.line.color.rgb = rgb('#FFFFFF'); badge.line.width = Pt(1.3)
    tf = badge.text_frame; tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    run = para.add_run(); run.text = str(operator_id)
    run.font.name = 'Arial'; run.font.bold = True
    run.font.size = Pt(8 if len(str(operator_id)) < 5 else 7)
    run.font.color.rgb = rgb('#FFFFFF')
    return (cx, cy)


def _copy_template_shapes(target_slide, source_slide):
    for shape in source_slide.shapes:
        target_slide.shapes._spTree.insert_element_before(deepcopy(shape._element), 'p:extLst')


def _load_or_create_sws_presentation():
    template = Path(__file__).resolve().parent / 'resources' / 'SWS_Frame.pptx'
    if template.is_file():
        prs = Presentation(str(template))
        # Keep a pristine copy of the frame for additional slides.
        source_prs = Presentation(str(template))
        return prs, source_prs.slides[0], True
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs, None, False


def _new_sws_slide(prs, template_slide=None):
    if template_slide is None:
        return prs.slides.add_slide(prs.slide_layouts[6])
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    _copy_template_shapes(slide, template_slide)
    return slide


def _table_shapes(slide):
    return [shape for shape in slide.shapes if getattr(shape, 'has_table', False)]


def sws_slide(prs, scenario, selected_operator=None, template_slide=None, reuse_existing_first=False):
    slide = prs.slides[0] if (reuse_existing_first and len(prs.slides) >= 1 and not selected_operator) else _new_sws_slide(prs, template_slide)
    all_operators = operator_ids(scenario)
    ids = [selected_operator] if selected_operator else all_operators

    # The uploaded SWS frame is a blank, user-editable form. Do not fill or
    # modify any of its existing metadata, tables, legend or footer fields.
    # Only draw the physical line and the selected operators' movement routes
    # in the empty central canvas.

    stations = scenario.get('stations', {})
    boxes = {k: v for k, v in scenario.get('boxes', {}).items() if v.get('type') in ('in', 'out', 'swip')}
    directions = [v for v in scenario.get('boxes', {}).values() if v.get('type') == 'direction']
    objects = [*stations.values(), *boxes.values()]
    for item in directions:
        angle = math.radians(float(item.get('rotation', 0) or 0))
        cx = float(item['x']) + float(item['w']) / 2; cy = float(item['y']) + float(item['h']) / 2
        rw = abs(float(item['w']) * math.cos(angle)) + abs(float(item['h']) * math.sin(angle))
        rh = abs(float(item['w']) * math.sin(angle)) + abs(float(item['h']) * math.cos(angle))
        objects.append({'x': cx - rw / 2, 'y': cy - rh / 2, 'w': rw, 'h': rh})
    if objects:
        x0 = min(float(v['x']) for v in objects); y0 = min(float(v['y']) for v in objects)
        x1 = max(float(v['x']) + float(v['w']) for v in objects)
        y1 = max(float(v['y']) + float(v['h']) for v in objects)
    else:
        x0 = y0 = 0.; x1 = 875.; y1 = 500.
    x0 -= 35.0; y0 -= 35.0; x1 += 35.0; y1 += 35.0
    world_w = max(250., x1 - x0); world_h = max(190., y1 - y0)

    # Use the supplied SWS frame exactly and draw inside its open body region.
    # Exact open body of the supplied frame, minus generous protection margins:
    # header ends at 1.178 in; legend starts at 6.19 in; footer at 6.70 in.
    BODY_X = 0.36; BODY_Y = 1.37; BODY_W = 12.60; BODY_H = 4.54
    scale = min(BODY_W / world_w, BODY_H / world_h)
    xpad = (BODY_W - world_w * scale) / 2
    ypad = (BODY_H - world_h * scale) / 2

    def xy(p):
        return (BODY_X + xpad + (float(p['x']) - x0) * scale,
                BODY_Y + ypad + (float(p['y']) - y0) * scale)

    arrow_png = Path(__file__).resolve().parent / 'static' / 'images' / 'line_direction.png'
    for direction in directions:
        x, y = xy(direction)
        w = max(.01, float(direction['w']) * scale)
        h = max(.01, float(direction['h']) * scale)
        if arrow_png.is_file():
            image = slide.shapes.add_picture(str(arrow_png), Inches(x), Inches(y), Inches(w), Inches(h))
            image.rotation = float(direction.get('rotation', 0) or 0)
            transform = image._element.spPr.xfrm
            if direction.get('flip_horizontal'): transform.set('flipH', '1')
            if direction.get('flip_vertical'): transform.set('flipV', '1')

    for sid, station in stations.items():
        x, y = xy(station)
        w = max(.36, station['w'] * scale); h = max(.25, station['h'] * scale)
        title = f"S{sid}. {station.get('name', 'Station')}"
        sws_box(slide, x, y, w, h, title[:55], bg='#ECF2F9', border='#426A92', size=7.0 if w < 1.3 else 8.3)
        if station.get('has_part'):
            add_part_icon(slide, x + w - .17, y + h - .19, size=.16)

    for bid, box in boxes.items():
        x, y = xy(box); w = max(.32, box['w'] * scale); h = max(.22, box['h'] * scale)
        kind = box.get('type')
        fill = {'in': '#E2F5E9', 'out': '#FFF0E9', 'swip': '#FFF6D9'}[kind]
        name = box.get('name', bid)
        label = 'SWIP' if kind == 'swip' else f"{kind.upper()}: {name}"
        if kind == 'in':
            label += f'\n×{box.get("start_count", 0)}'
        sws_box(slide, x, y, w, h, label[:56], bg=fill,
                border={'in': '#348C5B', 'out': '#C96D42', 'swip': '#BF941E'}[kind], size=7.3, dotted=True)
        if (kind == 'swip' and box.get('has_part')) or (kind == 'in' and int(box.get('start_count') or 0) > 0):
            add_part_icon(slide, x + w - .16, y + h - .18, size=.15)

    annotation_visits = []
    start_icons = []
    starts_per_location = {}
    for i, opid in enumerate(ids):
        operator = scenario['operators'][opid]
        hue = route_color(operator)
        visits = movement_visits(scenario, opid)
        if not visits:
            continue
        points = [xy({'x': visit['point'][0], 'y': visit['point'][1]}) for visit in visits]
        for j, current in enumerate(points):
            annotation_visits.append({'operator': opid, 'hue': hue,
                                      'point': current, 'number': j + 1})
        # This is a schematic, repeating operator work loop. No operator-home
        # or initial pickup arrow; connect each location and close the loop.
        if len(points) > 1:
            for j, start in enumerate(points):
                end = points[(j + 1) % len(points)]
                bend = .065 + (i % 4) * .03
                add_curved_path(slide, start, end, hue, bend=bend)
        start_loc = visits[0]['loc']
        offset_index = starts_per_location.get(start_loc, 0)
        starts_per_location[start_loc] = offset_index + 1
        start_icons.append((opid, hue, points[0], offset_index))

    # Draw the circular operator avatar at the starting location, then place
    # numbered circles around it without overlap.
    occupied = [add_operator_start_icon(slide, opid, hue, point, idx)
                for opid, hue, point, idx in start_icons]
    add_numbered_steps(slide, annotation_visits, reserved=occupied)



def export_sws_pptx(scenario):
    prs, template_slide, reuse_first = _load_or_create_sws_presentation()
    if len(prs.slides) == 0:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    sws_slide(prs, scenario, template_slide=template_slide, reuse_existing_first=reuse_first)
    for opid in operator_ids(scenario):
        sws_slide(prs, scenario, opid, template_slide=template_slide, reuse_existing_first=False)
    out = BytesIO(); prs.save(out); out.seek(0)
    return out


def swcs_style_cell(cell, *, bg=None, color='1C3653', bold=False, size=9, center=False, wrap=False):
    cell.font=Font(name='Arial',size=size,bold=bold,color=color.lstrip('#'))
    if bg:cell.fill=fill_rgb(bg)
    cell.alignment=Alignment(vertical='center',horizontal='center' if center else 'left',wrap_text=wrap)
    cell.border=Border(bottom=Side(style='hair',color='E1E8F1'))


def time_intersects(start,end,bucket_start,bucket_end):
    return end > bucket_start + 1e-8 and start < bucket_end - 1e-8


def make_swcs_sheet(ws, scenario, operator_id):
    rows=planned_steps(scenario, operator_id)
    takt=float(scenario.get('takt_time',30.0))
    total_work=sum(r['manual']+r['walk'] for r in rows)
    auto_horizon=max([total_work]+[r['auto_end'] for r in rows])
    horizon=max(takt,auto_horizon,1.0)
    bucket=horizon/42  # readable width, chosen per operator to fit one landscape page
    count=42
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='K10'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.print_options.horizontalCentered=True
    ws.page_margins=PageMargins(left=.25,right=.25,top=.34,bottom=.34,header=.1,footer=.1)
    ws.column_dimensions['A'].width=7
    for c in 'BCDEF':ws.column_dimensions[c].width=10
    for c in 'GHI':ws.column_dimensions[c].width=10
    ws.column_dimensions['J'].width=13
    for col in range(11,11+count):ws.column_dimensions[get_column_letter(col)].width=2.8
    last=get_column_letter(10+count)

    for rng in [f'A2:D2',f'A3:D3',f'E2:{last}3','A4:C4','D4:F4','G4:L4','M4:U4',f'V4:{last}4',
                'A5:C5','D5:F5','G5:L5','M5:U5',f'V5:{last}5',
                f'A7:{last}7','B8:F9','A8:A9','G8:G9','H8:H9','I8:I9','J8:J9']:
        ws.merge_cells(rng)
    headers={'A2':'Line Name','E2':'STANDARD WORK COMBINATION SHEET','A4':'Part / Frame','D4':str(scenario.get('framesize') or '—'),
             'G4':'Takt Time','M4':f'{takt:g} minutes','V4':f'Operator: {operator_id}',
             'A5':'Department','D5':str(scenario.get('department') or '—'),
             'G5':'Date','M5':date.today().strftime('%d / %b / %Y'),'V5':'DRAFT — Sequence-based',
             'A7':f'Operation Time Chart   •   1 division = {bucket:.2f} min   •   Manual / Automatic / Walking / Waiting',
             'A8':'Step','B8':'Operation Name / Instruction','G8':'Manual\n(min)','H8':'Auto\n(min)','I8':'Walk\n(min)','J8':'Walk\n(m)'}
    for cell,text in headers.items():
        ws[cell]=text
        swcs_style_cell(ws[cell],bg=BLUE if cell in ('E2','A7') else PALE,
                        color='FFFFFF' if cell in ('E2','A7') else BLUE,
                        bold=True,size=15 if cell=='E2' else 9,center=cell in ('E2','A7','A8','G8','H8','I8','J8'),wrap=True)
    ws['A3']=scenario_name(scenario);swcs_style_cell(ws['A3'],bold=True,size=11)
    for col in range(11,11+count):
        c=ws.cell(8,col); c.value=round((col-11)*bucket,1) if (col-11)%5==0 else None
        swcs_style_cell(c,bg=BLUE,color='FFFFFF',size=8,center=True)
        swcs_style_cell(ws.cell(9,col),bg='#E7EFF8',center=True)
    ws.row_dimensions[2].height=22;ws.row_dimensions[3].height=21
    ws.row_dimensions[7].height=23;ws.row_dimensions[8].height=21;ws.row_dimensions[9].height=20

    for n,step in enumerate(rows):
        row=10+n*2
        bottom=row+1
        ws.merge_cells(start_row=row,start_column=1,end_row=bottom,end_column=1)
        ws.merge_cells(start_row=row,start_column=2,end_row=bottom,end_column=6)
        for col,value in [(1,step['index']),(2,step['description']),(7,step['manual']),(8,step['auto']),
                          (9,step['walk']),(10,step['distance_m'])]:
            cell=ws.cell(row,col);cell.value=value
            swcs_style_cell(cell,bg='#F8FAFE' if n%2==0 else '#FFFFFF',
                            center=col!=2,wrap=col==2,bold=col==1)
            if col in (7,8,9,10):cell.number_format='0.00'
        for col in (7,8,9,10):
            swcs_style_cell(ws.cell(bottom,col),bg='#F8FAFE' if n%2==0 else '#FFFFFF')
        ws.row_dimensions[row].height=21
        ws.row_dimensions[bottom].height=16
        for col in range(11,11+count):
            t0=(col-11)*bucket;t1=t0+bucket
            # Editable DrawingML timing strokes are overlaid after save.
            # Keep the underlying cell grid subtle, like the supplied SWCS.
            swcs_style_cell(ws.cell(row,col),bg='#F1F5FB' if n%2==0 else '#FFFFFF',center=True)
            swcs_style_cell(ws.cell(bottom,col),bg='#E8EEF6' if n%2==0 else '#FAFBFE',center=True)
    total_row=10+2*len(rows)
    ws.merge_cells(start_row=total_row,start_column=1,end_row=total_row,end_column=6)
    ws.cell(total_row,1,'TOTAL');swcs_style_cell(ws.cell(total_row,1),bg=BLUE,color='FFFFFF',bold=True)
    for col in (7,8,9,10):
        cell=ws.cell(total_row,col)
        key={7:'manual',8:'auto',9:'walk',10:'distance_m'}[col]
        cell.value=sum(float(item[key]) for item in rows)
        cell.number_format='0.00';swcs_style_cell(cell,bg='#DEEAF6',bold=True,center=True)
    ws.merge_cells(start_row=total_row+2,start_column=1,end_row=total_row+2,end_column=5)
    ws.cell(total_row+2,1,'Symbols: MANUAL — solid →  |  WALKING ~ zigzag →  |  AUTOMATIC - - dashed →  |  WAITING not modeled')
    swcs_style_cell(ws.cell(total_row+2,1),color=BLUE,bold=True)
    for col,color in [(7,MANUAL),(9,AUTO),(11,WALK),(13,WAIT)]:
        for c in range(col,col+2):ws.cell(total_row+2,c).fill=fill_rgb(color)
    ws.merge_cells(start_row=total_row+4,start_column=1,end_row=total_row+4,end_column=10)
    ws.cell(total_row+4,1,'DRAFT: steps are derived from Pick / Work / Drop sequences. Auto time overlaps operator work; waiting and detailed safety / quality instructions require validation.')
    swcs_style_cell(ws.cell(total_row+4,1),color='8E4B1E',size=9,wrap=True)
    ws.row_dimensions[total_row+4].height=28
    ws.print_area=f'A2:{last}{total_row+4}'
    ws.oddFooter.center.text=f'{operator_id} | Standard Work Combination Sheet | Draft'


# Use editable native Excel DrawingML rather than a flattened chart image.
# The reference SWCS workbook also uses freeform shapes and lines.
XDR='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
DRAWING='http://schemas.openxmlformats.org/drawingml/2006/main'
SHEETML='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R_NS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG_NS='http://schemas.openxmlformats.org/package/2006/relationships'
CT_NS='http://schemas.openxmlformats.org/package/2006/content-types'
EMU=9525  # 1 CSS pixel
COL_PX=25  # approximate rendered width of a 2.8-character Excel column


def _tag(ns,tag):
    return f'{{{ns}}}{tag}'


def _time_anchor(t, bucket):
    # The time axis begins at K (zero-based column 10).
    v=max(0,float(t))/bucket
    whole=int(math.floor(v))
    part=v-whole
    return 10+whole,round(part*COL_PX*EMU)


def _shape_anchor(root, kind, start, end, row, bucket, color, shape_id, *,
                  stroke_height=0, flip_vertical=False, arrow=True):
    """Native editable Excel line shapes, including a segmented walk zigzag."""
    if end-start <= 1e-7: return
    if kind=='walking':
        # A zigzag is built from ordinary Excel line shapes. This is more
        # portable than custom freeforms and stays editable in Excel.
        px=max(12,round((end-start)/bucket*COL_PX))
        n=max(3,min(48,round(px/7)))
        full_time=max(end-start,12/COL_PX*bucket)
        for i in range(n):
            t0=start+i/n*full_time
            t1=start+(i+1)/n*full_time
            _shape_anchor(root,'walk-segment',t0,t1,row,bucket,color,shape_id+i,
                          stroke_height=8,flip_vertical=bool(i%2),arrow=(i==n-1))
        return
    col,off=_time_anchor(start,bucket)
    width=max(3,round((end-start)/bucket*COL_PX))
    height=max(0,int(stroke_height))
    container=ET.SubElement(root,_tag(XDR,'oneCellAnchor'))
    fr=ET.SubElement(container,_tag(XDR,'from'))
    for key,value in [('col',col),('colOff',off),('row',row),('rowOff',round(9*EMU))]:
        ET.SubElement(fr,_tag(XDR,key)).text=str(value)
    ET.SubElement(container,_tag(XDR,'ext'),{'cx':str(width*EMU),'cy':str(max(1,height)*EMU)})
    sp=ET.SubElement(container,_tag(XDR,'sp'))
    nv=ET.SubElement(sp,_tag(XDR,'nvSpPr'))
    ET.SubElement(nv,_tag(XDR,'cNvPr'),{'id':str(shape_id),'name':f'{kind.title()} timeline {shape_id}',
                                          'descr':f'{kind.title()} from {start:.3f} to {end:.3f} minutes'})
    ET.SubElement(nv,_tag(XDR,'cNvSpPr'))
    props=ET.SubElement(sp,_tag(XDR,'spPr'))
    xfrm=ET.SubElement(props,_tag(DRAWING,'xfrm'),{'flipV':'1'} if flip_vertical else {})
    ET.SubElement(xfrm,_tag(DRAWING,'off'),{'x':'0','y':'0'})
    ET.SubElement(xfrm,_tag(DRAWING,'ext'),{'cx':str(width*EMU),'cy':str(height*EMU)})
    geom=ET.SubElement(props,_tag(DRAWING,'prstGeom'),{'prst':'line'})
    ET.SubElement(geom,_tag(DRAWING,'avLst'))
    ET.SubElement(props,_tag(DRAWING,'noFill'))
    ln=ET.SubElement(props,_tag(DRAWING,'ln'),{'w':'24400'})
    sf=ET.SubElement(ln,_tag(DRAWING,'solidFill'))
    ET.SubElement(sf,_tag(DRAWING,'srgbClr'),{'val':color.lstrip('#')})
    if kind=='automatic': ET.SubElement(ln,_tag(DRAWING,'prstDash'),{'val':'dash'})
    ET.SubElement(ln,_tag(DRAWING,'round'))
    if arrow: ET.SubElement(ln,_tag(DRAWING,'tailEnd'),{'type':'triangle','w':'sm','len':'sm'})
    ET.SubElement(container,_tag(XDR,'clientData'))


def _drawing_layer_for_operator(scenario, operator_id):
    rows=planned_steps(scenario,operator_id)
    takt=float(scenario.get('takt_time',30.0))
    operator_total=sum(row['manual']+row['walk'] for row in rows)
    horizon=max(takt,operator_total,*(row['auto_end'] for row in rows),1.0)
    bucket=horizon/42
    root=ET.Element(_tag(XDR,'wsDr'))
    for i,step in enumerate(rows):
        index=10+i*2
        top=index-1  # zero-index row
        bottom=top+1
        _shape_anchor(root,'walking',step['walk_start'],step['manual_start'],top,bucket,WALK,1000+200*i+1)
        _shape_anchor(root,'manual',step['manual_start'],step['auto_start'],top,bucket,MANUAL,1000+200*i+150)
        _shape_anchor(root,'automatic',step['auto_start'],step['auto_end'],bottom,bucket,AUTO,1000+200*i+151)
    ET.register_namespace('xdr',XDR)
    ET.register_namespace('a',DRAWING)
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)


def _insert_native_drawings(workbook_bytes, scenario, ids):
    """Attach editable Excel line/freeform objects without changing cell data."""
    if not ids:return workbook_bytes
    src=BytesIO(workbook_bytes)
    dest=BytesIO()
    with ZipFile(src,'r') as inp, ZipFile(dest,'w',ZIP_DEFLATED) as out:
        archive={item.filename:inp.read(item.filename) for item in inp.infolist()}
        content=ET.fromstring(archive['[Content_Types].xml'])
        for idx,operator_id in enumerate(ids,1):
            sheet=f'xl/worksheets/sheet{idx}.xml'
            sheetroot=ET.fromstring(archive[sheet])
            # Standard worksheets have no drawings in the current exporter.
            if sheetroot.find(_tag(SHEETML,'drawing')) is not None:
                raise ValueError('SWCS sheet already has a drawing relationship')
            ET.SubElement(sheetroot,_tag(SHEETML,'drawing'),{_tag(R_NS,'id'):'rId1'})
            ET.register_namespace('',SHEETML)
            ET.register_namespace('r',R_NS)
            archive[sheet]=ET.tostring(sheetroot,encoding='utf-8',xml_declaration=True)
            rels=f'xl/worksheets/_rels/sheet{idx}.xml.rels'
            relationships=ET.Element(_tag(PKG_NS,'Relationships'))
            ET.SubElement(relationships,_tag(PKG_NS,'Relationship'),{'Id':'rId1',
                'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing',
                'Target':f'../drawings/drawing{idx}.xml'})
            ET.register_namespace('',PKG_NS)
            archive[rels]=ET.tostring(relationships,encoding='utf-8',xml_declaration=True)
            dname=f'xl/drawings/drawing{idx}.xml'
            archive[dname]=_drawing_layer_for_operator(scenario,operator_id)
            ET.SubElement(content,_tag(CT_NS,'Override'),{'PartName':'/'+dname,
                'ContentType':'application/vnd.openxmlformats-officedocument.drawing+xml'})
        ET.register_namespace('',CT_NS)
        archive['[Content_Types].xml']=ET.tostring(content,encoding='utf-8',xml_declaration=True)
        for path,data in archive.items():out.writestr(path,data)
    dest.seek(0)
    return dest


def _export_legacy_swcs_xlsx(scenario):
    wb=Workbook()
    ids=operator_ids(scenario)
    if not ids:
        ws=wb.active;ws.title='No operators'
        ws['A1']='No operators are configured for this line.'
    else:
        ws=wb.active;ws.title=ids[0][:31]
        make_swcs_sheet(ws,scenario,ids[0])
        for opid in ids[1:]:
            ws=wb.create_sheet(opid[:31]);make_swcs_sheet(ws,scenario,opid)
    out=BytesIO();wb.save(out);out.seek(0)
    return _insert_native_drawings(out.getvalue(),scenario,ids)


def export_swcs_xlsx(scenario):
    """Generate editable SWCS sheets using the actual supplied Excel template."""
    from swcs_preserved_template import export_template_swcs_xlsx
    return export_template_swcs_xlsx(scenario,operator_ids(scenario),planned_steps)
